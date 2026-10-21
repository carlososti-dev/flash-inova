"""Leitura da planilha da lista de autorizados (CSV ou XLSX), linha a linha, com os erros de cada uma."""
import csv
import io
import unicodedata
from dataclasses import dataclass

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook

from .models import Role

MAX_ROWS = 20000
TEMPLATE_CSV = (
    "matricula;nome;email;setor;perfil\r\n"
    "12345;Maria da Silva;maria.silva@hospital.org.br;Enfermagem;colaborador\r\n"
    "23456;João Souza;joao.souza@hospital.org.br;Qualidade;membro de comitê\r\n"
)

# Cabeçalhos aceitos (sem acento, minúsculos, só letras e números) → campo.
HEADERS = {
    "matricula": "matricula", "matr": "matricula", "registro": "matricula",
    "nome": "name", "nomecompleto": "name", "colaborador": "name",
    "email": "email", "emailcorporativo": "email",
    "setor": "sector", "area": "sector", "departamento": "sector",
    "perfil": "role", "papel": "role",
}
# Perfis que a planilha pode dar. Administrador fica de fora: é promovido na gestão de usuários.
ROLES = {
    "": None, "colaborador": Role.colaborador, "membrodecomite": Role.membro_comite, "membrocomite": Role.membro_comite,
    "comite": Role.membro_comite, "avaliador": Role.membro_comite, "gestor": Role.gestor,
}


class RosterFileError(ValueError):
    """Problema no arquivo inteiro (formato, cabeçalho, tamanho)."""


@dataclass
class RosterRow:
    line: int
    matricula: str
    name: str
    email: str
    sector: str | None
    role: Role | None


@dataclass
class RosterError:
    line: int
    message: str


def _key(value: str) -> str:
    plain = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return "".join(ch for ch in plain.lower() if ch.isalnum())


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # matrícula numérica no Excel chega como 12345.0
    return str(value).strip()


def _read_rows(data: bytes) -> list[list[str]]:
    if data.startswith(b"PK\x03\x04"):
        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # arquivo zip que não é planilha, ou corrompido
            raise RosterFileError("Não foi possível ler a planilha. Salve como .xlsx ou .csv e envie de novo.") from exc
        return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise RosterFileError("Planilha no formato antigo (.xls). Salve como .xlsx ou .csv e envie de novo.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")  # CSV salvo pelo Excel em português
    first = text.split("\n", 1)[0]
    delimiter = ";" if first.count(";") >= first.count(",") else ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def parse_roster(data: bytes) -> tuple[list[RosterRow], list[RosterError], int]:
    """Devolve (linhas válidas, erros, total de linhas com conteúdo). Duplicidades no próprio arquivo viram erro."""
    rows = _read_rows(data)
    while rows and not any(rows[0]):
        rows.pop(0)
    if not rows:
        raise RosterFileError("O arquivo está vazio.")
    columns: dict[str, int] = {}
    for index, header in enumerate(rows[0]):
        field = HEADERS.get(_key(header))
        if field and field not in columns:
            columns[field] = index
    missing = [label for field, label in (("matricula", "matricula"), ("name", "nome"), ("email", "email")) if field not in columns]
    if missing:
        raise RosterFileError(f"Faltam colunas no cabeçalho: {', '.join(missing)}. Use o modelo.")
    body = [(n, r) for n, r in enumerate(rows[1:], start=2) if any(r)]
    if len(body) > MAX_ROWS:
        raise RosterFileError(f"O arquivo tem mais de {MAX_ROWS} linhas. Divida em partes.")

    def get(row: list[str], field: str) -> str:
        i = columns.get(field)
        return row[i] if i is not None and i < len(row) else ""

    valid: list[RosterRow] = []
    errors: list[RosterError] = []
    seen_matricula: dict[str, int] = {}
    seen_email: dict[str, int] = {}
    for line, row in body:
        matricula, name, email = get(row, "matricula"), get(row, "name"), get(row, "email")
        problems = []
        if not matricula:
            problems.append("matrícula vazia")
        elif len(matricula) > 50:
            problems.append("matrícula longa demais")
        if not name:
            problems.append("nome vazio")
        if not email:
            problems.append("e-mail vazio")
        else:
            try:
                email = validate_email(email, check_deliverability=False).normalized.lower()
            except EmailNotValidError:
                problems.append(f"e-mail inválido ({email})")
        role_text = _key(get(row, "role"))
        if role_text not in ROLES:
            problems.append(f"perfil desconhecido ({get(row, 'role')}); use colaborador, membro de comitê ou gestor")
        if matricula in seen_matricula:
            problems.append(f"matrícula repetida (já está na linha {seen_matricula[matricula]})")
        if email and email in seen_email:
            problems.append(f"e-mail repetido (já está na linha {seen_email[email]})")
        if problems:
            errors.append(RosterError(line, "; ".join(problems).capitalize() + "."))
            continue
        seen_matricula[matricula] = line
        seen_email[email] = line
        valid.append(RosterRow(line, matricula, name[:200], email, get(row, "sector")[:200] or None, ROLES[role_text]))
    return valid, errors, len(body)
